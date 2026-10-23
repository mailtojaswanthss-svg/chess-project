from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the single-source Chesswise report.")
    parser.add_argument("metrics", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    payload = json.loads(args.metrics.read_text(encoding="utf-8"))
    metrics = payload["metrics"]
    label_definition = payload.get("label_definition", {})
    evaluation_split = payload.get("evaluation_split", {})
    document = Document()
    document.add_heading("Chesswise: Chess Coaching Case Study", 0)
    document.add_paragraph(
        "A local chess-training application combining Stockfish analysis, a calibrated "
        "blunder-risk ranking experiment, PGN review, trap mode, and an interactive web board."
    )

    document.add_heading("Results from one source of truth", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Validation result"
    table.rows[0].cells[2].text = "How to interpret it"
    rows = [
        ("PR-AUC", f"{metrics['pr_auc']:.3f}", "Primary ranking metric for this imbalanced task"),
        ("Top-10% lift", f"{metrics['top_decile_lift']:.2f}x", "Blunder rate among riskiest 10% divided by average rate"),
        ("ROC-AUC", f"{metrics['roc_auc']:.3f} ({metrics['roc_auc_ci_low']:.3f}-{metrics['roc_auc_ci_high']:.3f})", "Grouped game-bootstrap 95% interval"),
        ("Brier skill vs base rate", f"{metrics['brier_skill_vs_base_rate']:.1%}", "Improvement over always predicting the training base rate"),
        ("Brier score", f"{metrics['brier_score']:.3f}", "Lower is better; compare with the base-rate predictor"),
        ("Accuracy", f"{metrics['accuracy']:.1%} vs {metrics['majority_accuracy']:.1%}", "Footnote only; majority class is stronger"),
    ]
    for name, value, meaning in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, value, meaning

    document.add_heading("What the model can and cannot claim", level=1)
    document.add_paragraph(
        "The model is a position-ranking assistant, not an accurate standalone blunder detector. "
        "The data contains 1,038 games collected from four requested Lichess accounts, with 500 games "
        "each from two accounts; the games also contain many opponents. "
        f"Labels use Stockfish depth {label_definition.get('stockfish_depth', 8)} and a "
        f"{label_definition.get('loss_threshold_centipawns', 100)}-centipawn loss threshold. "
        "This is called an engine mistake label in this report because a one-pawn loss is not "
        "equivalent to a human-psychology blunder, and the labels are noisy. "
        "The confidence intervals are grouped by game."
    )
    if evaluation_split.get("strategy") == "player_held_out":
        document.add_paragraph(
            f"This run uses a player-held-out split across {len(evaluation_split.get('train_groups', []))} "
            f"training players and {len(evaluation_split.get('test_groups', []))} test players. "
            "The result is still preliminary because the source games are concentrated in four requested "
            "accounts and are not a representative sample of the Lichess population."
        )
    else:
        document.add_paragraph(
            "This run uses a game-held-out split. It does not measure generalization to new players."
        )
    document.add_paragraph(
        "No hard probability threshold is used in the product. The previous 0.348 threshold was "
        "removed after it produced very low recall on calibrated probabilities. The current product "
        "ranks positions and reports top-10% lift instead."
    )
    document.add_heading("Product and web experience", level=1)
    document.add_paragraph(
        "The web board supports coach-only play, Stockfish, trap mode, optional Maia-2, saved games, "
        "PGN export, and move-by-move review. Review cards identify whether White or Black made each "
        "move, show the position before that move, and separate engine-confirmed losses from predicted risk. "
        "The board now refreshes correctly after Undo and Reset, keeps the move list synchronized, and shows "
        "a current-position panel with the side to move and last move. Board controls and analysis cards are "
        "centered and responsive on smaller screens."
    )
    document.add_heading("Verified status", level=1)
    for item in [
        "11 backend tests pass.",
        "Frontend TypeScript/Vite production build passes.",
        "GET /health, /predict, /move, and /trap-move return successfully.",
        "Browser smoke check confirmed the updated board layout and current-position panel.",
        "Maia-2 weights and Docker Compose remain explicitly unverified environment work.",
    ]:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Next experiments", level=1)
    document.add_paragraph(
        "Collect more players across rating bands, use player-held-out evaluation, replace raw "
        "centipawn labels with win-probability loss, and validate ranking lift on a new player set."
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
